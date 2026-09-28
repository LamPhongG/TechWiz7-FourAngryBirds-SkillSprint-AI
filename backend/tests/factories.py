"""Builders for real PDF / DOCX files and API payloads used across tests."""
import io
import itertools
import uuid

import docx
import pymupdf

_codes = itertools.count(100)


def unique_code() -> str:
    """Codes outside the company catalog (DOC-01…28), so each test owns its document family."""
    return f"DOC-{next(_codes)}"


def make_pdf(pages: list[str]) -> bytes:
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = pdf.tobytes()
    pdf.close()
    return data


def make_scanned_pdf() -> bytes:
    """A page with only a drawing and no text layer, like a scanned document."""
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 300, 300), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    page.draw_line((60, 60), (290, 290))
    data = pdf.tobytes()
    pdf.close()
    return data


def make_docx(sections: list[tuple[str, list[str]]]) -> bytes:
    document = docx.Document()
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def policy_text(title: str = "Leave policy") -> list[str]:
    """Two PDF pages of policy text; a random marker keeps every file's SHA-256 unique."""
    marker = uuid.uuid4().hex[:8]
    return [
        f"{title.upper()}\n1. Scope\nThis policy applies to all employees. Reference {marker}.\n"
        "2. Annual leave\nEmployees must request annual leave at least 5 days in advance.",
        "3. Sick leave\nEmployees should notify their manager before 9 AM on the day of absence.",
    ]


def upload(client, headers, content: bytes, file_name: str, **meta):
    # Default category is "FAQ" (optional), not "Policy": since Coverage Score (Pipeline 2),
    # all "Policy"/"Company-wide" documents are considered mandatory for all onboarding paths during
    # the test session (shared DB, session scope). Tests that require a mandatory document should pass category explicitly.
    fields = {
        "code": meta.pop("code", None) or unique_code(),
        "title_en": meta.pop("title_en", None) or f"Policy {uuid.uuid4().hex[:6]}",
        "category": "FAQ",
        "department_code": "Company-wide",
        "version": "1.0",
        "effective_date": "2026-01-01",
    } | meta
    return client.post("/api/documents", headers=headers, data=fields, files={"file": (file_name, content)})


def mandatory_source_ids(client, headers, position_id: str) -> list[str]:
    """Ready mandatory documents of the matrix for a position: a client must send them as sources, as the UI does.

    Other tests upload catalog documents (DOC-10…) into the shared test database, which makes them mandatory."""
    res = client.get(f"/api/job-positions/{position_id}/required-sources", headers=headers)
    assert res.status_code == 200, res.text
    return [s["document"]["id"] for s in res.json() if s["mandatory"] and s["status"] == "ready"]


def upload_ready_pdf(client, headers, **meta) -> dict:
    res = upload(client, headers, make_pdf(policy_text()), "policy.pdf", **meta)
    assert res.status_code == 201, res.text
    assert res.json()["processing_status"] == "ready"
    return res.json()


def path_content(doc: dict, chunks: list[dict], prefix: str = "LP-TEST", stage: str = "day1") -> dict:
    """Minimal generated content citing real chunks of `doc`, shaped like the frontend generator's."""
    chunk = chunks[0]
    quote = chunk["content"][:60]
    ref = {"doc_id": doc["id"], "doc": doc["code"], "section": chunk["heading"], "page": chunk["page"],
           "chunk_id": chunk["chunk_id"], "exact_quote": quote}
    # The correct option must appear inside the quote, or the knowledge check calls it a contradiction.
    correct = max(quote.split(), key=len)
    module_id = f"{prefix}-M1"
    return {
        "stages": [{
            "key": stage,
            "modules": [{
                "id": module_id, "kind": "lesson", "title": doc["title"], "titleEn": doc["title_en"],
                "doc_id": doc["id"], "doc_code": doc["code"], "tier": 1,
                "lessons": [{"id": f"{module_id}-L1", "title": chunk["heading"], "content": chunk["content"],
                             "minutes": 1, "source_reference": ref}],
                "tasks": [{"id": f"{module_id}-T1", "title": "Request leave 5 days ahead",
                           "completion_criteria": "Leave request submitted in the HR portal 5 days ahead", "source_reference": ref}],
                "quiz": [{"id": f"{module_id}-Q1", "kind": "statement", "question": "Which is correct?",
                          "questionEn": "Which is correct?", "options": ["Zebra", correct, "Quokka"], "answer": 1,
                          "source_reference": ref}],
            }],
        }],
        "excluded_chunks": [],
        "coverage": None,
        "engine": "local-draft",
        "model": None,
        "prompt_version": "v1.0",
    }
